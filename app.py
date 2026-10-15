"""ATS Resume Checker - Streamlit + Gemini Flash."""
import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_CHARS = 20000  # keep prompts small and fast

PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Evaluate the resume below{jd_clause}. Be honest, specific and strict.

Return ONLY valid JSON with exactly this schema:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "breakdown": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "content_quality": <0-100>,
    "structure": <0-100>,
    "impact_metrics": <0-100>
  }},
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"section": "<section name>", "issue": "<problem>", "suggestion": "<how to fix>"}}
  ],
  "rewrite_examples": [
    {{"before": "<original bullet from resume>", "after": "<improved bullet>"}}
  ]
}}

Rules: 3-6 items per list, 2-4 rewrite examples taken from the actual resume.
Never invent experience the candidate does not have.

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def parse_json(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if not match:
            raise ValueError("The AI returned an unreadable response. Please retry.")
        return json.loads(match.group(0))


def clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(float(value))))
    except (TypeError, ValueError):
        return default


def analyze_resume(client, resume_text: str, job_desc: str = "") -> dict:
    jd_clause = " against the job description provided" if job_desc.strip() else ""
    jd_block = f'JOB DESCRIPTION:\n"""\n{job_desc[:MAX_CHARS]}\n"""\n\n' if job_desc.strip() else ""
    prompt = PROMPT.format(jd_clause=jd_clause, jd_block=jd_block, resume=resume_text[:MAX_CHARS])
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json", temperature=0.2
        ),
    )
    result = parse_json(response.text)
    result["ats_score"] = clamp(result.get("ats_score"))
    result["breakdown"] = {k: clamp(v) for k, v in (result.get("breakdown") or {}).items()}
    return result


def get_api_key() -> str:
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY", "")


def render(result: dict) -> None:
    score = result["ats_score"]
    label = "Excellent" if score >= 80 else "Good" if score >= 65 else "Needs work" if score >= 50 else "Poor"
    c1, c2 = st.columns([1, 2])
    c1.metric("ATS Score", f"{score}/100", label)
    c2.write(result.get("summary", ""))
    st.progress(score / 100)

    if result["breakdown"]:
        st.subheader("Score breakdown")
        cols = st.columns(len(result["breakdown"]))
        for col, (k, v) in zip(cols, result["breakdown"].items()):
            col.metric(k.replace("_", " ").title(), v)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result.get("strengths", []):
            st.markdown(f"- {s}")
    with right:
        st.subheader("Weaknesses")
        for w in result.get("weaknesses", []):
            st.markdown(f"- {w}")

    if result.get("missing_keywords"):
        st.subheader("Missing keywords")
        st.write(" ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("Suggested improvements")
    for item in result.get("improvements", []):
        with st.expander(f"{item.get('section', 'General')}: {item.get('issue', '')}"):
            st.write(item.get("suggestion", ""))

    if result.get("rewrite_examples"):
        st.subheader("Rewrite examples")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex.get('before', '')}")
            st.markdown(f"**After:** {ex.get('after', '')}")
            st.divider()


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it. Powered by Gemini.")

    api_key = get_api_key()
    if not api_key:
        api_key = st.sidebar.text_input("Gemini API key", type="password",
                                        help="Get a free key at aistudio.google.com")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_desc = st.text_area("Job description (optional, improves keyword matching)", height=150)

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key.")
            return
        try:
            text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return
        if len(text) < 50:
            st.error("Almost no text was found. If your resume is a scanned image, "
                     "ATS systems can't read it either. Export a text-based PDF.")
            return
        with st.spinner("Analyzing your resume..."):
            try:
                result = analyze_resume(genai.Client(api_key=api_key), text, job_desc)
            except Exception as e:
                st.error(f"Analysis failed: {e}")
                return
        render(result)


if __name__ == "__main__":
    main()
